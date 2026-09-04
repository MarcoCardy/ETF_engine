from __future__ import annotations

import csv
import hashlib
import io
import math
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from typing import Mapping, Sequence

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from perpetual_engine.data_sources import SourceArtifact
from perpetual_engine.io import canonical_json
from perpetual_engine.point_in_time import ObservationRow, validate_rows


CANONICAL_WORLD_LABEL = "PUBLIC_DEVELOPED_WORLD_TR_PROXY_USD"
_INTERNATIONAL_URL = "https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp/F-F_International_Indices.zip"
_US_URL = "https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp/F-F_Research_Data_Factors_CSV.zip"
_DEVELOPED_URL = "https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp/Developed_3_Factors_CSV.zip"
_MOMENTUM_URL = "https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp/Developed_6_Portfolios_ME_Prior_12_2_CSV.zip"
_QUALITY_URL = "https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp/Developed_6_Portfolios_ME_OP_CSV.zip"
_AQR_TSMOM_URL = "https://www.aqr.com/-/media/AQR/Documents/Insights/Data-Sets/Time-Series-Momentum-Factors-Monthly.xlsx"
_WDI_ROOT = "https://api.worldbank.org/v2"
_SENTINELS = {Decimal("-99.99"), Decimal("-999")}
_COMPOSITE_SOURCE_URL = "derived://public-developed-world/monthly-return"


@dataclass(frozen=True)
class PublicWorldInputs:
    international_archive: SourceArtifact
    us_archive: SourceArtifact
    developed_archive: SourceArtifact
    wdi_caps: Mapping[str, Mapping[int, Decimal]]
    wdi_source_url: str
    wdi_source_hash: str
    retrieved_at: datetime


@dataclass(frozen=True)
class InputProvenance:
    source_url: str
    retrieved_at: datetime
    source_hash: str


@dataclass(frozen=True)
class PublicWorldReturn:
    observation: ObservationRow
    segment: str
    input_provenance: tuple[InputProvenance, ...]


@dataclass(frozen=True)
class FXQuote:
    observation: ObservationRow
    source_observation: ObservationRow
    raw: float
    raw_quote: str = "USD_per_EUR"
    segment: str = "DEXUSEU"

    @property
    def month(self) -> date:
        return self.observation.observation_date

    @property
    def eur_per_usd(self) -> float:
        return float(self.observation.value)


@dataclass(frozen=True)
class FXDailyCarry:
    observation_date: date
    fixing: FXQuote

    @property
    def value(self) -> float:
        return self.fixing.eur_per_usd


@dataclass(frozen=True)
class ProxyMetrics:
    n: int
    correlation: float
    beta: float
    tracking_error: float
    cagr_gap: float
    max_drawdown_gap: float
    status: str


def _month_start(value: date) -> date:
    return value.replace(day=1)


def _month_end(value: date) -> date:
    month = _month_start(value)
    return _next_month(month) - timedelta(days=1)


def _next_month(value: date) -> date:
    return date(value.year + (value.month == 12), value.month % 12 + 1, 1)


def _decimal(value: str, field: str) -> Decimal:
    try:
        result = Decimal(value.strip())
    except (InvalidOperation, AttributeError) as error:
        raise ValueError(f"{field} must be numeric") from error
    if not result.is_finite() or result in _SENTINELS:
        raise ValueError(f"{field} is missing or non-finite")
    return result


def _month(value: str) -> date:
    value = value.strip()
    if not re.fullmatch(r"\d{6}", value):
        raise ValueError("French monthly date must be YYYYMM")
    return _month_end(date(int(value[:4]), int(value[4:]), 1))


def _require_artifact(artifact: SourceArtifact, expected_url: str) -> None:
    if artifact.source_url != expected_url:
        raise ValueError(f"source URL must be exactly {expected_url}")
    if artifact.retrieved_at.tzinfo is None or artifact.retrieved_at.utcoffset() is None:
        raise ValueError("artifact retrieved_at must be timezone-aware")


def _artifact_provenance(artifact: SourceArtifact) -> InputProvenance:
    return InputProvenance(artifact.source_url, artifact.retrieved_at.astimezone(timezone.utc), artifact.source_hash)


def _ordered_provenance(sources: Sequence[InputProvenance]) -> tuple[InputProvenance, ...]:
    return tuple(sorted(sources, key=lambda source: (source.source_url, source.retrieved_at, source.source_hash)))


def _composite_source_hash(sources: Sequence[InputProvenance]) -> str:
    return hashlib.sha256(
        canonical_json(
            [
                {"source_hash": source.source_hash, "source_url": source.source_url, "retrieved_at": source.retrieved_at.isoformat()}
                for source in _ordered_provenance(sources)
            ]
        )
    ).hexdigest()


def _zip_member(artifact: SourceArtifact, member: str) -> str:
    _require_artifact(artifact, {
        "Ind_all.Dat": _INTERNATIONAL_URL,
        "F-F_Research_Data_Factors.csv": _US_URL,
        "Developed_3_Factors.csv": _DEVELOPED_URL,
        "Developed_6_Portfolios_ME_Prior_12_2.csv": _MOMENTUM_URL,
        "Developed_6_Portfolios_ME_OP.csv": _QUALITY_URL,
    }[member])
    try:
        with zipfile.ZipFile(artifact.local_path) as archive:
            if member not in archive.namelist():
                raise ValueError(f"archive must contain exact member {member}")
            return archive.read(member).decode("utf-8-sig")
    except zipfile.BadZipFile as error:
        raise ValueError("French artifact must be a ZIP archive") from error


def _international_returns(artifact: SourceArtifact) -> dict[date, Decimal]:
    lines = _zip_member(artifact, "Ind_all.Dat").splitlines()
    try:
        start = next(index for index, line in enumerate(lines) if line.strip() == "Value-Weight Dollar Returns")
    except StopIteration as error:
        raise ValueError("international archive requires exact Value-Weight Dollar Returns table") from error
    header_index = next((index for index in range(start + 1, len(lines)) if lines[index].strip()), None)
    if header_index is None:
        raise ValueError("international archive has no table header")
    headers = lines[header_index].split()
    if headers.count("Mkt") != 1:
        raise ValueError("international archive requires exact Mkt column")
    date_index = headers.index("Date") if headers.count("Date") == 1 else 0
    market_index = headers.index("Mkt")
    result: dict[date, Decimal] = {}
    for line in lines[header_index + 1 :]:
        cells = line.split()
        if not cells or not re.fullmatch(r"\d{6}", cells[0]):
            continue
        if max(date_index, market_index) >= len(cells):
            raise ValueError("international archive has a malformed data row")
        month = _month(cells[date_index])
        if month in result:
            raise ValueError("international archive has duplicate months")
        result[month] = _decimal(cells[market_index], "Mkt") / Decimal("100")
    if not result:
        raise ValueError("international archive has no monthly Mkt returns")
    return result


def _factor_returns(artifact: SourceArtifact, member: str, *, risk_free_only: bool = False) -> dict[date, Decimal]:
    text = _zip_member(artifact, member)
    lines = text.splitlines()
    header_index = next((index for index, line in enumerate(lines) if "Mkt-RF" in line and "RF" in line), None)
    if header_index is None:
        raise ValueError("French factor archive requires Mkt-RF and RF columns")
    reader = csv.DictReader(io.StringIO("\n".join(lines[header_index:])))
    if reader.fieldnames is None or reader.fieldnames.count("Mkt-RF") != 1 or reader.fieldnames.count("RF") != 1:
        raise ValueError("French factor archive requires exact Mkt-RF and RF columns")
    date_field = reader.fieldnames[0]
    result: dict[date, Decimal] = {}
    for row in reader:
        raw_date = (row.get(date_field) or "").strip()
        if not raw_date:
            continue
        if not re.fullmatch(r"\d{6}", raw_date):
            continue
        month = _month(raw_date)
        if month in result:
            raise ValueError("French factor archive has duplicate months")
        market_excess = _decimal(row["Mkt-RF"] or "", "Mkt-RF")
        risk_free = _decimal(row["RF"] or "", "RF")
        result[month] = (risk_free if risk_free_only else market_excess + risk_free) / Decimal("100")
    if not result:
        raise ValueError("French factor archive has no monthly returns")
    return result


def _style_returns(
    artifact: SourceArtifact,
    member: str,
    expected_columns: tuple[str, ...],
    selected_column: str,
) -> dict[date, Decimal]:
    lines = _zip_member(artifact, member).splitlines()
    try:
        table_index = next(index for index, line in enumerate(lines) if line.strip() == "Average Value Weighted Returns -- Monthly")
    except StopIteration as error:
        raise ValueError("French style archive requires exact monthly value-weighted table") from error
    header_index = next((index for index in range(table_index + 1, len(lines)) if lines[index].strip()), None)
    if header_index is None:
        raise ValueError("French style archive has no table header")
    headers = tuple(cell.strip() for cell in next(csv.reader((lines[header_index],))))
    if headers != ("", *expected_columns):
        raise ValueError("French style archive columns do not match the exact audited table")
    try:
        end_index = next(
            index
            for index in range(header_index + 1, len(lines))
            if lines[index].strip() == "Average Equal Weighted Returns -- Monthly"
        )
    except StopIteration as error:
        raise ValueError("French style archive is truncated before the next audited table") from error
    selected_index = headers.index(selected_column)
    result: dict[date, Decimal] = {}
    for line in lines[header_index + 1 : end_index]:
        if not line.strip():
            continue
        cells = tuple(cell.strip() for cell in next(csv.reader((line,))))
        if len(cells) != len(headers) or not re.fullmatch(r"\d{6}", cells[0]):
            raise ValueError("French style archive has a malformed row after monthly data begins")
        month = _month(cells[0])
        if month in result:
            raise ValueError("French style archive has duplicate months")
        result[month] = _decimal(cells[selected_index], selected_column) / Decimal("100")
    if not result:
        raise ValueError("French style archive has no monthly returns")
    _require_continuous_months(tuple(sorted(result)), "French style archive")
    return result


def _blank(value: object) -> bool:
    return value is None or isinstance(value, str) and not value.strip()


def _aqr_tsmom_returns(artifact: SourceArtifact) -> dict[date, Decimal]:
    _require_artifact(artifact, _AQR_TSMOM_URL)
    try:
        workbook = load_workbook(io.BytesIO(artifact.local_path.read_bytes()), read_only=True, data_only=True)
    except (InvalidFileException, OSError, zipfile.BadZipFile) as error:
        raise ValueError("AQR TSMOM artifact must be an XLSX workbook") from error
    if "TSMOM Factors" not in workbook.sheetnames:
        raise ValueError("AQR TSMOM workbook requires exact TSMOM Factors sheet")
    sheet = workbook["TSMOM Factors"]
    title = sheet["A1"].value
    if not isinstance(title, str) or not title.endswith("Time Series Momentum: Factors, Monthly"):
        raise ValueError("AQR TSMOM workbook requires the exact monthly factor table")
    if sheet["A3"].value != "This file contains the excess returns of the long/short Time Series Momentum (TSMOM) factors.":
        raise ValueError("AQR TSMOM workbook requires the excess-return unit declaration")
    headers = tuple(sheet.cell(18, column).value for column in range(1, 7))
    if headers != (None, "TSMOM", "TSMOM^CM", "TSMOM^EQ", "TSMOM^FI", "TSMOM^FX"):
        raise ValueError("AQR TSMOM workbook columns do not match the exact audited table")
    data_rows = list(sheet.iter_rows(min_row=19, max_col=6))
    dated_rows = [index for index, row in enumerate(data_rows) if isinstance(row[0].value, (date, datetime))]
    if not dated_rows:
        raise ValueError("AQR TSMOM workbook has no monthly returns")
    last_dated_row = dated_rows[-1]
    result: dict[date, Decimal] = {}
    for row in data_rows[: last_dated_row + 1]:
        raw_date = row[0].value
        if not isinstance(raw_date, (date, datetime)):
            raise ValueError("AQR TSMOM workbook has a malformed row after monthly data begins")
        values = tuple(cell.value for cell in row[1:6])
        if any(not isinstance(value, (int, float, Decimal)) or isinstance(value, bool) for value in values):
            raise ValueError("AQR TSMOM workbook has a malformed factor value")
        if any(cell.number_format != "0.00%" for cell in row[1:6]):
            raise ValueError("AQR TSMOM workbook factor units must be decimal returns formatted as percent")
        month = _month_end(raw_date.date() if isinstance(raw_date, datetime) else raw_date)
        if month in result:
            raise ValueError("AQR TSMOM workbook has duplicate months")
        result[month] = _decimal(str(values[0]), "TSMOM")
    for row in data_rows[last_dated_row + 1 :]:
        if any(not _blank(cell.value) for cell in row):
            raise ValueError("AQR TSMOM workbook has trailing malformed data")
    _require_continuous_months(tuple(sorted(result)), "AQR TSMOM workbook")
    return result


def _require_continuous_months(months: Sequence[date], label: str) -> None:
    for previous, current in zip(months, months[1:]):
        if _next_month(_month_start(previous)) != _month_start(current):
            raise ValueError(f"{label} has an internal month gap")


def parse_french_archive(artifact: SourceArtifact, parser_kind: str) -> tuple[ObservationRow, ...]:
    """Parse exact audited academic source contracts into normalized observations."""
    contracts = {
        "french_international_archive": (
            lambda: _international_returns(artifact),
            "FRENCH_INTERNATIONAL_MKT_USD",
            "FRENCH_INTERNATIONAL_VALUE_WEIGHT_MKT",
        ),
        "french_us_factor_archive": (
            lambda: _factor_returns(artifact, "F-F_Research_Data_Factors.csv"),
            "FRENCH_US_MKT_USD",
            "FRENCH_US_MKT_RF_PLUS_RF",
        ),
        "french_developed_factor_archive": (
            lambda: _factor_returns(artifact, "Developed_3_Factors.csv"),
            "FRENCH_DEVELOPED_MKT_USD",
            "FRENCH_DEVELOPED_MKT_RF_PLUS_RF",
        ),
        "french_developed_rf_archive": (
            lambda: _factor_returns(artifact, "Developed_3_Factors.csv", risk_free_only=True),
            "FRENCH_DEVELOPED_RF_USD",
            "FRENCH_DEVELOPED_RF",
        ),
        "french_developed_momentum_archive": (
            lambda: _style_returns(
                artifact,
                "Developed_6_Portfolios_ME_Prior_12_2.csv",
                ("SMALL LoPRIOR", "ME1 PRIOR2", "SMALL HiPRIOR", "BIG LoPRIOR", "ME2 PRIOR2", "BIG HiPRIOR"),
                "BIG HiPRIOR",
            ),
            "FRENCH_DEVELOPED_BIG_HIPRIOR_USD",
            "FRENCH_VALUE_WEIGHTED_BIG_HIPRIOR",
        ),
        "french_developed_quality_archive": (
            lambda: _style_returns(
                artifact,
                "Developed_6_Portfolios_ME_OP.csv",
                ("SMALL LoOP", "ME1 OP2", "SMALL HiOP", "BIG LoOP", "ME2 OP2", "BIG HiOP"),
                "BIG HiOP",
            ),
            "FRENCH_DEVELOPED_BIG_ROBUST_USD",
            "FRENCH_VALUE_WEIGHTED_BIG_ROBUST_FROM_BIG_HIOP",
        ),
        "french_aqr_tsmom_archive": (
            lambda: _aqr_tsmom_returns(artifact),
            "AQR_TSMOM_EXCESS_RETURN_USD",
            "AQR_DIVERSIFIED_TSMOM_EXCESS_RETURN",
        ),
    }
    if parser_kind not in contracts:
        raise ValueError("unknown audited French parser kind")
    parse, series_id, flag = contracts[parser_kind]
    return validate_rows(
        ObservationRow(
            series_id=series_id,
            observation_date=month,
            period_end=month,
            available_at=datetime(month.year, month.month, month.day, 23, 59, 59, tzinfo=timezone.utc),
            value=value,
            unit="ratio",
            source_url=artifact.source_url,
            retrieved_at=artifact.retrieved_at,
            source_hash=artifact.source_hash,
            quality_flags=(flag,),
        )
        for month, value in sorted(parse().items())
    )


def _continuous_months(months: Sequence[date]) -> None:
    _require_continuous_months(months, "public world proxy")


def _wdi_weight(caps: Mapping[str, Mapping[int, Decimal]], year: int) -> Decimal:
    cap_year = year - 2
    try:
        usa, world = caps["USA"][cap_year], caps["WLD"][cap_year]
    except KeyError as error:
        raise ValueError(f"WDI USA and WLD values are mandatory for {cap_year}") from error
    if not isinstance(usa, Decimal) or not isinstance(world, Decimal) or not usa.is_finite() or not world.is_finite() or world <= 0:
        raise ValueError("WDI caps must be finite Decimals and World cap positive")
    weight = usa / world
    if not Decimal("0") <= weight <= Decimal("1"):
        raise ValueError("WDI USA weight must be within [0,1]")
    return weight


def build_public_world_monthly(inputs: PublicWorldInputs) -> tuple[PublicWorldReturn, ...]:
    """Build the frozen public developed-world USD monthly return proxy offline."""
    if inputs.wdi_source_url != _WDI_ROOT or not re.fullmatch(r"[0-9a-f]{64}", inputs.wdi_source_hash):
        raise ValueError("WDI source must use the exact API root and a SHA-256 hash")
    if inputs.retrieved_at.tzinfo is None or inputs.retrieved_at.utcoffset() is None:
        raise ValueError("retrieved_at must be timezone-aware")
    international = _international_returns(inputs.international_archive)
    us = _factor_returns(inputs.us_archive, "F-F_Research_Data_Factors.csv")
    developed = _factor_returns(inputs.developed_archive, "Developed_3_Factors.csv")
    reconstructed_months = sorted(month for month in international if date(1977, 1, 31) <= month <= date(1990, 6, 30))
    if not reconstructed_months or reconstructed_months[0] != date(1977, 1, 31):
        raise ValueError("reconstructed warm-up must begin January 1977")
    if any(month not in us for month in reconstructed_months):
        raise ValueError("US and international returns must both exist for every reconstructed month")
    if date(1990, 7, 31) not in developed:
        raise ValueError("developed return splice must begin July 1990")
    output: list[PublicWorldReturn] = []
    wdi = InputProvenance(inputs.wdi_source_url, inputs.retrieved_at.astimezone(timezone.utc), inputs.wdi_source_hash)
    reconstructed_sources = (_artifact_provenance(inputs.international_archive), _artifact_provenance(inputs.us_archive), wdi)
    developed_sources = (_artifact_provenance(inputs.developed_archive),)
    for month in reconstructed_months:
        weight = _wdi_weight(inputs.wdi_caps, month.year)
        value = weight * us[month] + (Decimal("1") - weight) * international[month]
        output.append(_public_row(month, value, "RECONSTRUCTED_SPLICE", reconstructed_sources))
    for month in sorted(month for month in developed if month >= date(1990, 7, 31)):
        output.append(_public_row(month, developed[month], "FF_DEVELOPED", developed_sources))
    months = [item.observation.observation_date for item in output]
    _continuous_months(months)
    return tuple(output)


def _public_row(month: date, value: Decimal, segment: str, sources: Sequence[InputProvenance]) -> PublicWorldReturn:
    provenance = _ordered_provenance(sources)
    row = ObservationRow(
        CANONICAL_WORLD_LABEL,
        month,
        month,
        datetime(month.year, month.month, month.day, 23, 59, 59, tzinfo=timezone.utc),
        value,
        "ratio",
        _COMPOSITE_SOURCE_URL,
        max(source.retrieved_at for source in provenance),
        _composite_source_hash(provenance),
        "CURRENT_VINTAGE_RESEARCH",
        (segment, "WDI_Y_MINUS_2", "COMPOSITE_PROVENANCE") if segment == "RECONSTRUCTED_SPLICE" else (segment, "COMPOSITE_PROVENANCE"),
    )
    return PublicWorldReturn(validate_rows((row,))[0], segment, provenance)


def convert_usd_to_eur(return_usd: float, fx_prev: float, fx_t: float) -> float:
    if not all(math.isfinite(value) for value in (return_usd, fx_prev, fx_t)) or fx_prev <= 0 or fx_t <= 0:
        raise ValueError("USD return and EUR_per_USD FX values must be finite, with FX positive")
    return (1.0 + return_usd) * (fx_t / fx_prev) - 1.0


def normalize_fx_monthly(rows: Sequence[ObservationRow]) -> tuple[FXQuote, ...]:
    """Select each final non-null raw USD-per-EUR fixing and normalize to EUR-per-USD."""
    selected: dict[date, ObservationRow] = {}
    for row in validate_rows(rows):
        if row.series_id not in {"CCUSSP01DEM650N", "DEXUSEU"}:
            raise ValueError("FX series must be CCUSSP01DEM650N or DEXUSEU")
        if row.unit != "USD_per_EUR" or row.value <= 0:
            raise ValueError("FX raw quote must be positive USD_per_EUR")
        month = _month_end(row.observation_date)
        previous = selected.get(month)
        if previous is None or _pit_key(row) > _pit_key(previous):
            selected[month] = row
    output: list[FXQuote] = []
    for month, row in sorted(selected.items()):
        if row.series_id == "CCUSSP01DEM650N":
            if not date(1977, 1, 31) <= month <= date(1998, 12, 31):
                raise ValueError("synthetic EUR FX must be January 1977 through December 1998")
            segment = "SYNTHETIC_EUR_FX"
        else:
            if month < date(1999, 1, 31):
                raise ValueError("DEXUSEU must begin in 1999")
            segment = "DEXUSEU"
        raw = float(row.value)
        normalized = 1.0 / raw
        if month == date(2023, 12, 31) and not (1.095 <= raw <= 1.115 and 0.900 <= normalized <= 0.910):
            raise ValueError("December 2023 DEXUSEU quote-direction sentinel failed")
        normalized_row = ObservationRow(
            f"EUR_PER_USD_{row.series_id}",
            month,
            month,
            row.available_at,
            Decimal(str(normalized)),
            "EUR_per_USD",
            row.source_url,
            row.retrieved_at,
            row.source_hash,
            row.vintage_status,
            tuple((*row.quality_flags, segment, "NORMALIZED_FROM_USD_PER_EUR")),
        )
        output.append(FXQuote(validate_rows((normalized_row,))[0], row, raw, "USD_per_EUR", segment))
    return tuple(output)


def _pit_key(row: ObservationRow) -> tuple[date, datetime, datetime, str]:
    return (row.observation_date, row.available_at, row.retrieved_at, row.source_hash)


def carry_fx_daily(fixings: Sequence[FXQuote], underlying_dates: Sequence[date]) -> dict[date, FXDailyCarry]:
    if any(not math.isfinite(quote.eur_per_usd) or quote.eur_per_usd <= 0 for quote in fixings):
        raise ValueError("FX fixings must be positive finite EUR_per_USD values")
    by_day = {quote.source_observation.observation_date: quote for quote in fixings}
    if len(by_day) != len(fixings):
        raise ValueError("FX fixings must have unique source fixing dates")
    ordered = sorted(by_day)
    output: dict[date, FXDailyCarry] = {}
    for day in underlying_dates:
        prior = next((candidate for candidate in reversed(ordered) if candidate <= day), None)
        if prior is None or (day - prior).days > 7:
            raise ValueError("FX carry exceeds seven calendar days")
        output[day] = FXDailyCarry(day, by_day[prior])
    return output


def _return_mapping(values: Mapping[date, float] | Sequence[PublicWorldReturn]) -> dict[date, float]:
    if isinstance(values, Mapping):
        pairs = ((_month_end(month), float(value)) for month, value in values.items())
    else:
        pairs = ((value.observation.observation_date, float(value.observation.value)) for value in values)
    result = dict(pairs)
    if len(result) != len(values) or any(not isinstance(month, date) or not math.isfinite(value) or value < -1 for month, value in result.items()):
        raise ValueError("return series must have unique finite dated returns")
    return result


def _empty_metrics(status: str) -> ProxyMetrics:
    return ProxyMetrics(0, math.nan, math.nan, math.nan, math.nan, math.nan, status)


def _cagr(values: Sequence[float]) -> float:
    wealth = math.prod(1.0 + value for value in values)
    return wealth ** (12.0 / len(values)) - 1.0


def _max_drawdown(values: Sequence[float]) -> float:
    wealth = peak = 1.0
    drawdown = 0.0
    for value in values:
        wealth *= 1.0 + value
        peak = max(peak, wealth)
        drawdown = min(drawdown, wealth / peak - 1.0)
    return drawdown


def validate_world_proxy(
    proxy_eur_returns: Mapping[date, float] | Sequence[PublicWorldReturn],
    official_nav_eur_returns: Mapping[date, float] | None,
) -> ProxyMetrics:
    if official_nav_eur_returns is None:
        return _empty_metrics("FAIL_MISSING_OFFICIAL_NAV")
    try:
        proxy, nav = _return_mapping(proxy_eur_returns), _return_mapping(official_nav_eur_returns)
    except (TypeError, ValueError):
        return _empty_metrics("FAIL_INSUFFICIENT_NAV_SAMPLE")
    if not proxy or not nav:
        return _empty_metrics("FAIL_INSUFFICIENT_NAV_SAMPLE")
    common_start, common_end = max(min(proxy), min(nav)), min(max(proxy), max(nav))
    if common_start > common_end:
        return _empty_metrics("FAIL_INSUFFICIENT_NAV_SAMPLE")
    months: list[date] = []
    current = _month_start(common_start)
    while current <= _month_start(common_end):
        month = _month_end(current)
        if month not in proxy or month not in nav:
            return _empty_metrics("FAIL_INSUFFICIENT_NAV_SAMPLE")
        months.append(month)
        current = _next_month(current)
    if len(months) < 60:
        return _empty_metrics("FAIL_INSUFFICIENT_NAV_SAMPLE")
    p, n = [proxy[month] for month in months], [nav[month] for month in months]
    mean_p, mean_n = sum(p) / len(p), sum(n) / len(n)
    variance_n = sum((value - mean_n) ** 2 for value in n) / (len(n) - 1)
    variance_p = sum((value - mean_p) ** 2 for value in p) / (len(p) - 1)
    if variance_n == 0 or variance_p == 0:
        return _empty_metrics("FAIL_INSUFFICIENT_NAV_SAMPLE")
    covariance = sum((a - mean_p) * (b - mean_n) for a, b in zip(p, n)) / (len(p) - 1)
    correlation = covariance / math.sqrt(variance_p * variance_n)
    beta = covariance / variance_n
    differences = [a - b for a, b in zip(p, n)]
    tracking_error = math.sqrt(sum((value - sum(differences) / len(differences)) ** 2 for value in differences) / (len(differences) - 1)) * math.sqrt(12)
    cagr_gap = abs(_cagr(p) - _cagr(n))
    max_drawdown_gap = abs(_max_drawdown(p) - _max_drawdown(n))
    status = "PASS" if correlation >= 0.98 and 0.95 <= beta <= 1.05 and tracking_error <= 0.03 and cagr_gap <= 0.015 else "FAIL"
    return ProxyMetrics(len(months), correlation, beta, tracking_error, cagr_gap, max_drawdown_gap, status)
