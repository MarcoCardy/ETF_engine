from __future__ import annotations

import hashlib
import json
import ssl
import urllib.request
import csv
from dataclasses import dataclass, replace
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import unquote, urlparse

import xlrd
import certifi
from openpyxl import load_workbook

from perpetual_engine.io import canonical_json, sha256_file
from perpetual_engine.point_in_time import ObservationRow, asof_select, validate_rows


_RESERVED_ARTIFACT_NAMES = frozenset({"manifest.json", "provenance", "raw"})
_WINDOWS_INVALID_FILENAME_CHARS = frozenset('<>:"/\\|?*')
_WINDOWS_DEVICE_NAMES = frozenset(
    {"con", "prn", "aux", "nul", *(f"com{number}" for number in range(1, 10)), *(f"lpt{number}" for number in range(1, 10))}
)

_DAMODARAN_ANNUAL_URL = "https://pages.stern.nyu.edu/~adamodar/pc/datasets/histimpl.xls"
_DAMODARAN_MONTHLY_URL = "https://pages.stern.nyu.edu/~adamodar/pc/implprem/ERPbymonth.xlsx"
_DAMODARAN_ANNUAL_FIELD = "Implied ERP (FCFE)"
_DAMODARAN_MONTHLY_SHEET = "Historical ERP"
_DAMODARAN_MONTHLY_DATE_FIELD = "Start of month"
_DAMODARAN_MONTHLY_VALUE_FIELD = "ERP (T12m)"
_DAMODARAN_MONTHLY_START = date(2008, 9, 1)
TREASURY_SIGNAL_IDS = frozenset({"DGS10", "DFII10"})
DEFENSIVE_RATE_ID = "IR3TIB01ITM156N"
_FRED_ROOT = "https://fred.stlouisfed.org/graph/fredgraph.csv?id="
_WDI_MARKET_CAP_URL = "https://api.worldbank.org/v2/country/USA;WLD/indicator/CM.MKT.LCAP.CD?format=json&per_page=20000"
_EUROSTAT_HICP_URL = "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/prc_hicp_midx?geo=IT&coicop=CP00&unit=I15"


def fetch_url(url: str) -> bytes:
    """The sole network adapter; parsers and backtests remain offline."""
    if not isinstance(url, str) or not url.startswith("https://"):
        raise ValueError("network source URL must be HTTPS")
    with urllib.request.urlopen(url, context=ssl.create_default_context(cafile=certifi.where())) as response:
        return response.read()


def _windows_collision_key(filename: str) -> str:
    return filename.rstrip(" .").casefold()


def _validated_windows_filename_key(filename: str) -> str:
    if any(character in _WINDOWS_INVALID_FILENAME_CHARS or ord(character) < 32 for character in filename):
        raise ValueError("source_url contains an invalid filename character")
    collision_key = _windows_collision_key(filename)
    if not collision_key:
        raise ValueError("source_url must include a usable filename")
    device_name = collision_key.split(".", 1)[0].rstrip()
    if device_name in _WINDOWS_DEVICE_NAMES:
        raise ValueError("source_url filename is a reserved device name")
    return collision_key


@dataclass(frozen=True)
class SourceArtifact:
    source_url: str
    retrieved_at: datetime
    source_hash: str
    local_path: Path
    byte_count: int
    parser_version: str


def _require_artifact_integrity(artifact: SourceArtifact, expected_url: str) -> None:
    if not isinstance(artifact, SourceArtifact) or artifact.source_url != expected_url:
        raise ValueError(f"source URL must be exactly {expected_url}")
    if artifact.retrieved_at.tzinfo is None or artifact.retrieved_at.utcoffset() is None:
        raise ValueError("artifact retrieved_at must be timezone-aware")
    if not artifact.local_path.is_file() or sha256_file(artifact.local_path) != artifact.source_hash:
        raise ValueError("source artifact SHA-256 does not reconcile")


def _month_end(value: date) -> date:
    following = date(value.year + (value.month == 12), value.month % 12 + 1, 1)
    return following - timedelta(days=1)


def parse_fred_csv(
    artifact: SourceArtifact,
    *,
    series_id: str,
    unit: str,
    source_scale: str,
    frequency: str,
    availability: str,
) -> tuple[ObservationRow, ...]:
    """Parse one explicitly configured official FRED CSV without inferring semantics."""
    _require_artifact_integrity(artifact, _FRED_ROOT + series_id)
    if source_scale not in {"raw", "percent"} or frequency not in {"daily", "monthly"}:
        raise ValueError("FRED source scale/frequency is invalid")
    with artifact.local_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        date_field = next((field for field in ("observation_date", "DATE") if field in (reader.fieldnames or ())), None)
        if date_field is None or tuple(reader.fieldnames or ()).count(series_id) != 1 or len(reader.fieldnames or ()) != 2:
            raise ValueError(f"FRED CSV requires exact date and {series_id} columns")
        output: list[ObservationRow] = []
        seen: set[date] = set()
        for record in reader:
            raw = (record.get(series_id) or "").strip()
            if raw in {"", "."}:
                continue
            try:
                observed = date.fromisoformat((record.get(date_field) or "").strip())
                value = Decimal(raw)
            except (InvalidOperation, ValueError) as error:
                raise ValueError(f"FRED {series_id} row is malformed") from error
            if not value.is_finite():
                raise ValueError(f"FRED {series_id} value must be finite")
            normalized_date = _month_end(observed) if frequency == "monthly" else observed
            if normalized_date in seen:
                raise ValueError(f"FRED {series_id} has duplicate observations")
            seen.add(normalized_date)
            if availability == "treasury_next_business_day":
                available_at = treasury_available_at(observed)
            elif availability == "ir3_m_plus_1":
                available_at = ir3tib_available_at(observed)
            elif availability == "month_end":
                available_at = datetime(
                    normalized_date.year, normalized_date.month, normalized_date.day, 23, 59, 59, tzinfo=timezone.utc
                )
            elif availability == "observation_end":
                available_at = datetime(observed.year, observed.month, observed.day, 23, 59, 59, tzinfo=timezone.utc)
            else:
                raise ValueError(f"FRED {series_id} availability rule is invalid")
            if available_at > artifact.retrieved_at:
                continue
            output.append(
                ObservationRow(
                    series_id=series_id,
                    observation_date=normalized_date,
                    period_end=normalized_date,
                    available_at=available_at,
                    value=value / Decimal("100") if source_scale == "percent" else value,
                    unit=unit,
                    source_url=artifact.source_url,
                    retrieved_at=artifact.retrieved_at,
                    source_hash=artifact.source_hash,
                    quality_flags=("FRED_OFFICIAL", availability.upper()),
                )
            )
    if not output:
        raise ValueError(f"FRED {series_id} has no usable observations")
    return validate_rows(output)


def parse_wdi_market_cap_json(
    artifact: SourceArtifact,
    *,
    indicator: str,
    aggregates: tuple[str, ...],
) -> tuple[ObservationRow, ...]:
    _require_artifact_integrity(artifact, _WDI_MARKET_CAP_URL)
    if indicator != "CM.MKT.LCAP.CD" or aggregates != ("USA", "WLD"):
        raise ValueError("WDI requires exact CM.MKT.LCAP.CD USA/WLD contract")
    try:
        payload = json.loads(artifact.local_path.read_text(encoding="utf-8"))
        records = payload[1]
    except (IndexError, TypeError, json.JSONDecodeError) as error:
        raise ValueError("WDI response must be the official JSON array") from error
    if not isinstance(records, list):
        raise ValueError("WDI response records must be an array")
    output: list[ObservationRow] = []
    identities: set[tuple[str, int]] = set()
    for record in records:
        if not isinstance(record, dict) or not isinstance(record.get("indicator"), dict):
            raise ValueError("WDI record is malformed")
        if record["indicator"].get("id") != indicator:
            raise ValueError("WDI indicator must be CM.MKT.LCAP.CD")
        aggregate = record.get("countryiso3code")
        if aggregate not in aggregates:
            raise ValueError("WDI aggregate must be USA or WLD")
        try:
            year = int(record["date"])
            value = Decimal(str(record["value"]))
        except (InvalidOperation, KeyError, TypeError, ValueError) as error:
            raise ValueError("WDI year/value is malformed") from error
        identity = (aggregate, year)
        if identity in identities or not value.is_finite() or value <= 0:
            raise ValueError("WDI values must be unique, finite, and positive")
        identities.add(identity)
        observed = date(year, 12, 31)
        output.append(
            ObservationRow(
                f"WDI_{indicator}_{aggregate}", observed, observed, artifact.retrieved_at, value, "current_USD",
                artifact.source_url, artifact.retrieved_at, artifact.source_hash, quality_flags=("WDI_OFFICIAL", aggregate),
            )
        )
    years = {year for _, year in identities}
    if not output or any((aggregate, year) not in identities for aggregate in aggregates for year in years):
        raise ValueError("WDI requires paired USA/WLD values for every year")
    return validate_rows(sorted(output, key=lambda row: (row.observation_date, row.series_id)))


def parse_eurostat_hicp_json(artifact: SourceArtifact) -> tuple[ObservationRow, ...]:
    _require_artifact_integrity(artifact, _EUROSTAT_HICP_URL)
    try:
        payload = json.loads(artifact.local_path.read_text(encoding="utf-8"))
        dimensions = payload["dimension"]
        ids, sizes = payload["id"], payload["size"]
        values = payload["value"]
    except (KeyError, TypeError, json.JSONDecodeError) as error:
        raise ValueError("Eurostat HICP response is malformed") from error
    expected_ids = ["freq", "unit", "coicop", "geo", "time"]
    if ids != expected_ids or sizes[:4] != [1, 1, 1, 1] or len(sizes) != 5:
        raise ValueError("Eurostat HICP dimensions must be exact")
    expected_dimensions = {"freq": "M", "unit": "I15", "coicop": "CP00", "geo": "IT"}
    for dimension, expected in expected_dimensions.items():
        try:
            index = dimensions[dimension]["category"]["index"]
        except (KeyError, TypeError) as error:
            raise ValueError(f"Eurostat {dimension} dimension is malformed") from error
        if index != {expected: 0}:
            raise ValueError(f"Eurostat {dimension} dimension must be {expected}")
    try:
        time_index = dimensions["time"]["category"]["index"]
    except (KeyError, TypeError) as error:
        raise ValueError("Eurostat time dimension is malformed") from error
    if not isinstance(time_index, dict) or len(time_index) != sizes[4]:
        raise ValueError("Eurostat time dimension size does not reconcile")
    output: list[ObservationRow] = []
    for period, index in sorted(time_index.items(), key=lambda item: item[1]):
        try:
            month_start = date.fromisoformat(f"{period}-01")
            raw_value = values.get(str(index)) if isinstance(values, dict) else values[index]
            value = Decimal(str(raw_value))
        except (AttributeError, IndexError, InvalidOperation, TypeError, ValueError) as error:
            raise ValueError("Eurostat HICP time/value is malformed") from error
        if raw_value is None or not value.is_finite() or value <= 0:
            raise ValueError("Eurostat HICP value must be finite and positive")
        observed = _month_end(month_start)
        output.append(
            ObservationRow(
                "HICP", observed, observed, artifact.retrieved_at, value, "I15_index", artifact.source_url,
                artifact.retrieved_at, artifact.source_hash, quality_flags=("EUROSTAT_OFFICIAL", "ITALY_CP00"),
            )
        )
    if not output:
        raise ValueError("Eurostat HICP has no observations")
    return validate_rows(output)


def _require_damodaran_artifact(artifact: SourceArtifact, expected_url: str) -> None:
    if artifact.source_url != expected_url:
        raise ValueError(f"Damodaran source URL must be exactly {expected_url}")
    if artifact.retrieved_at.tzinfo is None or artifact.retrieved_at.utcoffset() is None:
        raise ValueError("artifact retrieved_at must be timezone-aware")


def _decimal_cell(value: object, field: str) -> Decimal:
    if value is None or isinstance(value, bool):
        raise ValueError(f"{field} must be a non-null numeric value")
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError) as error:
        raise ValueError(f"{field} must be a non-null numeric value") from error
    if not result.is_finite():
        raise ValueError(f"{field} must be a finite numeric value")
    return result


def _normalize_erp(value: object, field: str, percent_formatted: bool) -> Decimal:
    raw = _decimal_cell(value, field)
    if raw < 0:
        raise ValueError(f"{field} cannot be negative")
    if raw >= Decimal("1"):
        return raw / Decimal("100")
    if percent_formatted:
        return raw
    raise ValueError(f"{field} decimal scale requires a percent-formatted workbook cell")


def _validate_erp_scale(rows: tuple[ObservationRow, ...]) -> tuple[ObservationRow, ...]:
    values = sorted(row.value for row in rows)
    middle = len(values) // 2
    median = values[middle] if len(values) % 2 else (values[middle - 1] + values[middle]) / Decimal("2")
    if median < Decimal("0.01") or median > Decimal("0.15"):
        raise ValueError("normalized ERP historical median must be between 1% and 15%")
    return rows


def _annual_available_at(year: int) -> datetime:
    return datetime(year + 1, 2, 1, 23, 59, 59, tzinfo=timezone.utc)


def _observed_fixed_holiday(year: int, month: int, day: int) -> date:
    holiday = date(year, month, day)
    if holiday.weekday() == 5:
        return holiday - timedelta(days=1)
    if holiday.weekday() == 6:
        return holiday + timedelta(days=1)
    return holiday


def _nth_weekday(year: int, month: int, weekday: int, occurrence: int) -> date:
    result = date(year, month, 1)
    result += timedelta(days=(weekday - result.weekday()) % 7 + (occurrence - 1) * 7)
    return result


def _last_weekday(year: int, month: int, weekday: int) -> date:
    if month == 12:
        result = date(year + 1, 1, 1) - timedelta(days=1)
    else:
        result = date(year, month + 1, 1) - timedelta(days=1)
    return result - timedelta(days=(result.weekday() - weekday) % 7)


def _us_federal_holidays(year: int) -> frozenset[date]:
    holidays = {
        _observed_fixed_holiday(year, 1, 1),
        _nth_weekday(year, 1, 0, 3),
        _nth_weekday(year, 2, 0, 3),
        _last_weekday(year, 5, 0),
        _observed_fixed_holiday(year, 7, 4),
        _nth_weekday(year, 9, 0, 1),
        _nth_weekday(year, 10, 0, 2),
        _observed_fixed_holiday(year, 11, 11),
        _nth_weekday(year, 11, 3, 4),
        _observed_fixed_holiday(year, 12, 25),
    }
    if year >= 2021:
        holidays.add(_observed_fixed_holiday(year, 6, 19))
    return frozenset(holidays)


def _first_business_day(month: date) -> date:
    result = month.replace(day=1)
    holidays = _us_federal_holidays(result.year)
    while result.weekday() >= 5 or result in holidays:
        result += timedelta(days=1)
    return result


def treasury_available_at(observation_date: date, archived_release_at: datetime | None = None) -> datetime:
    """Return the official timestamp or next-US-business-day PIT fallback."""
    if archived_release_at is not None:
        if archived_release_at.tzinfo is None or archived_release_at.utcoffset() is None:
            raise ValueError("archived_release_at must be timezone-aware")
        return archived_release_at.astimezone(timezone.utc)
    if not isinstance(observation_date, date) or isinstance(observation_date, datetime):
        raise ValueError("observation_date must be a date")
    result = observation_date + timedelta(days=1)
    while result.weekday() >= 5 or result in _us_federal_holidays(result.year):
        result += timedelta(days=1)
    return datetime(result.year, result.month, result.day, 23, 59, 59, tzinfo=timezone.utc)


def ir3tib_available_at(observation_month: date, archived_release_at: datetime | None = None) -> datetime:
    """Return the official timestamp or conservative M+1 final-instant fallback."""
    if archived_release_at is not None:
        if archived_release_at.tzinfo is None or archived_release_at.utcoffset() is None:
            raise ValueError("archived_release_at must be timezone-aware")
        return archived_release_at.astimezone(timezone.utc)
    if not isinstance(observation_month, date) or isinstance(observation_month, datetime):
        raise ValueError("observation_month must be a date")
    next_month = date(observation_month.year + (observation_month.month == 12), observation_month.month % 12 + 1, 1)
    final_day = date(next_month.year + (next_month.month == 12), next_month.month % 12 + 1, 1) - timedelta(days=1)
    return datetime(final_day.year, final_day.month, final_day.day, 23, 59, 59, tzinfo=timezone.utc)


def rate_available_at(series_id: str, observation_date: date, archived_release_at: datetime | None = None) -> datetime:
    """Dispatch Task 5 timing without a generic guessed release lag."""
    if series_id in TREASURY_SIGNAL_IDS:
        return treasury_available_at(observation_date, archived_release_at)
    if series_id == DEFENSIVE_RATE_ID:
        return ir3tib_available_at(observation_date, archived_release_at)
    raise ValueError(f"no Task 5 availability rule for {series_id}")


def normalize_rate_availability(row: ObservationRow, archived_release_at: datetime | None = None) -> ObservationRow:
    """Apply the exact Task 5 fallback to a frozen normalized rate observation."""
    available_at = rate_available_at(row.series_id, row.observation_date, archived_release_at)
    if row.retrieved_at.astimezone(timezone.utc) < available_at:
        raise ValueError("retrieved_at cannot precede normalized rate availability")
    return validate_rows((replace(row, available_at=available_at),))[0]


def _monthly_available_at(month: date) -> datetime:
    business_day = _first_business_day(month)
    return datetime(business_day.year, business_day.month, business_day.day, 23, 59, 59, tzinfo=timezone.utc)


def _load_damodaran_monthly_max_staleness_days() -> int:
    config_path = Path(__file__).resolve().parents[1] / "config" / "data_sources_v1.json"
    try:
        config = json.loads(config_path.read_text(encoding="utf-8"))
        days = config["damodaran"]["monthly"]["max_staleness_days"]
    except (OSError, KeyError, TypeError, json.JSONDecodeError) as error:
        raise ValueError("Damodaran monthly freshness configuration is invalid") from error
    if not isinstance(days, int) or isinstance(days, bool) or days < 0:
        raise ValueError("Damodaran monthly max_staleness_days must be a non-negative integer")
    return days


def _validate_monthly_freshness(rows: tuple[ObservationRow, ...], artifact: SourceArtifact) -> None:
    latest = max(row.observation_date for row in rows)
    age_days = (artifact.retrieved_at.astimezone(timezone.utc).date() - latest).days
    if age_days > _load_damodaran_monthly_max_staleness_days():
        raise ValueError("Damodaran monthly maximum observation date is stale")


def _annual_percent_format(workbook: xlrd.book.Book, cell: xlrd.sheet.Cell) -> bool:
    try:
        format_key = workbook.xf_list[cell.xf_index].format_key
        return "%" in workbook.format_map[format_key].format_str
    except (AttributeError, IndexError, KeyError, TypeError):
        return False


def _year(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and float(value).is_integer():
        return int(value)
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    return None


def parse_damodaran_annual(path: Path, artifact: SourceArtifact) -> tuple[ObservationRow, ...]:
    _require_damodaran_artifact(artifact, _DAMODARAN_ANNUAL_URL)
    workbook = xlrd.open_workbook(str(path), formatting_info=True)
    header: tuple[xlrd.sheet.Sheet, int, int, int] | None = None
    for sheet in workbook.sheets():
        for row_index in range(sheet.nrows):
            cells = [str(cell.value).strip() for cell in sheet.row(row_index)]
            if "Year" in cells and _DAMODARAN_ANNUAL_FIELD in cells:
                if header is not None:
                    raise ValueError("Damodaran annual workbook has ambiguous headers")
                header = (sheet, row_index, cells.index("Year"), cells.index(_DAMODARAN_ANNUAL_FIELD))
    if header is None:
        raise ValueError(f"Damodaran annual workbook requires Year and {_DAMODARAN_ANNUAL_FIELD}")
    sheet, header_row, year_column, value_column = header
    rows: list[ObservationRow] = []
    dates: set[date] = set()
    for row_index in range(header_row + 1, sheet.nrows):
        year = _year(sheet.cell_value(row_index, year_column))
        if year is None:
            continue
        if year > 2007:
            continue
        value_cell = sheet.cell(row_index, value_column)
        if value_cell.value is None or str(value_cell.value).strip() == "":
            continue
        observation_date = date(year, 12, 31)
        if observation_date in dates:
            raise ValueError("Damodaran annual workbook has duplicate observation dates")
        dates.add(observation_date)
        rows.append(
            ObservationRow(
                series_id="DAMODARAN_ERP_T12M",
                observation_date=observation_date,
                period_end=observation_date,
                available_at=_annual_available_at(year),
                value=_normalize_erp(value_cell.value, _DAMODARAN_ANNUAL_FIELD, _annual_percent_format(workbook, value_cell)),
                unit="ratio",
                source_url=artifact.source_url,
                retrieved_at=artifact.retrieved_at,
                source_hash=artifact.source_hash,
            )
        )
    if not rows:
        raise ValueError("Damodaran annual workbook has no usable observations")
    years = sorted(row.observation_date.year for row in rows)
    if any(current != previous + 1 for previous, current in zip(years, years[1:])):
        raise ValueError("Damodaran annual workbook has an internal year gap")
    if years[-1] != 2007:
        raise ValueError("Damodaran annual workbook must end in 2007")
    return _validate_erp_scale(validate_rows(rows))


def _monthly_columns(sheet: object) -> tuple[int, int, int]:
    for row_index, row in enumerate(sheet.iter_rows()):  # type: ignore[attr-defined]
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in row]
        if _DAMODARAN_MONTHLY_DATE_FIELD in headers and _DAMODARAN_MONTHLY_VALUE_FIELD in headers:
            if headers.count(_DAMODARAN_MONTHLY_DATE_FIELD) != 1 or headers.count(_DAMODARAN_MONTHLY_VALUE_FIELD) != 1:
                raise ValueError("Damodaran monthly workbook has ambiguous required columns")
            return row_index, headers.index(_DAMODARAN_MONTHLY_DATE_FIELD), headers.index(_DAMODARAN_MONTHLY_VALUE_FIELD)
    raise ValueError(
        f"Damodaran monthly workbook requires {_DAMODARAN_MONTHLY_DATE_FIELD} and {_DAMODARAN_MONTHLY_VALUE_FIELD}"
    )


def parse_damodaran_monthly(path: Path, artifact: SourceArtifact) -> tuple[ObservationRow, ...]:
    _require_damodaran_artifact(artifact, _DAMODARAN_MONTHLY_URL)
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[ObservationRow] = []
    dates: set[date] = set()
    try:
        if _DAMODARAN_MONTHLY_SHEET not in workbook.sheetnames:
            raise ValueError(f"Damodaran monthly workbook requires sheet {_DAMODARAN_MONTHLY_SHEET}")
        sheet = workbook[_DAMODARAN_MONTHLY_SHEET]
        header_row, date_column, value_column = _monthly_columns(sheet)
        for row in sheet.iter_rows(min_row=header_row + 2):
            date_cell = row[date_column]
            value_cell = row[value_column]
            if date_cell.value is None and value_cell.value is None:
                continue
            if isinstance(date_cell.value, datetime):
                observation_date = date_cell.value.date()
            elif isinstance(date_cell.value, date):
                observation_date = date_cell.value
            elif isinstance(date_cell.value, str):
                try:
                    observation_date = datetime.strptime(date_cell.value.strip(), "%d-%b-%y").date()
                except ValueError as error:
                    raise ValueError(f"{_DAMODARAN_MONTHLY_DATE_FIELD} must be a date") from error
            else:
                raise ValueError(f"{_DAMODARAN_MONTHLY_DATE_FIELD} must be a date")
            if observation_date.day != 1 or observation_date < _DAMODARAN_MONTHLY_START:
                raise ValueError("Damodaran monthly observations must begin in September 2008 at month start")
            if observation_date in dates:
                raise ValueError("Damodaran monthly workbook has duplicate observation dates")
            dates.add(observation_date)
            number_format = str(value_cell.number_format or "")
            value = value_cell.value
            if isinstance(value, str) and value.strip().endswith("%"):
                value = value.strip()[:-1]
            rows.append(
                ObservationRow(
                    series_id="DAMODARAN_ERP_T12M",
                    observation_date=observation_date,
                    period_end=observation_date,
                    available_at=_monthly_available_at(observation_date),
                    value=_normalize_erp(value, _DAMODARAN_MONTHLY_VALUE_FIELD, "%" in number_format),
                    unit="ratio",
                    source_url=artifact.source_url,
                    retrieved_at=artifact.retrieved_at,
                    source_hash=artifact.source_hash,
                )
            )
    finally:
        workbook.close()
    if not rows:
        raise ValueError("Damodaran monthly workbook has no usable observations")
    if min(dates) != _DAMODARAN_MONTHLY_START:
        raise ValueError("Damodaran monthly workbook first observation must be 2008-09-01")
    ordered_dates = sorted(dates)
    for previous, current in zip(ordered_dates, ordered_dates[1:]):
        expected = date(previous.year + (previous.month == 12), previous.month % 12 + 1, 1)
        if current != expected:
            raise ValueError("Damodaran monthly workbook has an internal date gap")
    validated = _validate_erp_scale(validate_rows(rows))
    _validate_monthly_freshness(validated, artifact)
    return validated


def select_damodaran_erp(
    annual_rows: tuple[ObservationRow, ...], monthly_rows: tuple[ObservationRow, ...], decision_at: datetime
) -> ObservationRow | None:
    if decision_at.tzinfo is None or decision_at.utcoffset() is None:
        raise ValueError("decision_at must be timezone-aware")
    switch_at = datetime(2008, 9, 1, tzinfo=timezone.utc)
    if decision_at.astimezone(timezone.utc) >= switch_at:
        return asof_select((row for row in monthly_rows if row.observation_date >= _DAMODARAN_MONTHLY_START), decision_at)
    return asof_select((row for row in annual_rows if row.observation_date < _DAMODARAN_MONTHLY_START), decision_at)


def damodaran_erp_asof(
    annual_rows: tuple[ObservationRow, ...], monthly_rows: tuple[ObservationRow, ...], decision_at: datetime
) -> ObservationRow | None:
    return select_damodaran_erp(annual_rows, monthly_rows, decision_at)


def freeze_bytes(
    content: bytes,
    source_url: str,
    retrieved_at: datetime,
    raw_dir: Path,
    parser_version: str,
) -> SourceArtifact:
    if retrieved_at.tzinfo is None or retrieved_at.utcoffset() is None:
        raise ValueError("retrieved_at must be timezone-aware")
    filename = unquote(urlparse(source_url).path.rsplit("/", 1)[-1])
    if not filename:
        raise ValueError("source_url must include a filename")
    collision_key = _validated_windows_filename_key(filename)
    reserved_name = collision_key in _RESERVED_ARTIFACT_NAMES
    source_hash = hashlib.sha256(content).hexdigest()
    artifact_dir = raw_dir / source_hash
    local_path = (
        artifact_dir / "raw" / f"{hashlib.sha256(source_url.encode()).hexdigest()}.bin"
        if reserved_name
        else artifact_dir / filename
    )
    manifest_path = artifact_dir / "manifest.json"
    retrieved_at = retrieved_at.astimezone(timezone.utc)
    record = {
        "byte_count": len(content),
        "parser_version": parser_version,
        "retrieved_at": retrieved_at.isoformat(),
        "source_hash": source_hash,
        "source_url": source_url,
    }
    record_bytes = canonical_json(record)
    record_hash = hashlib.sha256(record_bytes).hexdigest()
    provenance_path = artifact_dir / "provenance" / f"{record_hash}.json"
    artifact_dir.mkdir(parents=True, exist_ok=True)
    local_path.parent.mkdir(exist_ok=True)
    if local_path.exists():
        if sha256_file(local_path) != source_hash:
            raise ValueError(f"corrupt frozen artifact: {local_path}")
    else:
        local_path.write_bytes(content)
    if not manifest_path.exists():
        manifest_path.write_bytes(record_bytes)
    provenance_path.parent.mkdir(exist_ok=True)
    if not provenance_path.exists():
        provenance_path.write_bytes(record_bytes)
    return SourceArtifact(
        source_url=source_url,
        retrieved_at=retrieved_at,
        source_hash=source_hash,
        local_path=local_path,
        byte_count=len(content),
        parser_version=parser_version,
    )
