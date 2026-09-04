from __future__ import annotations

import io
import math
import zipfile
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from openpyxl import Workbook

from perpetual_engine.data_sources import SourceArtifact
from perpetual_engine.market_proxy import (
    CANONICAL_WORLD_LABEL,
    FXQuote,
    PublicWorldInputs,
    build_public_world_monthly,
    carry_fx_daily,
    convert_usd_to_eur,
    normalize_fx_monthly,
    parse_french_archive,
    validate_world_proxy,
)
from perpetual_engine.point_in_time import ObservationRow, asof_select


NOW = datetime(2026, 8, 22, tzinfo=timezone.utc)
MOMENTUM_URL = "https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp/Developed_6_Portfolios_ME_Prior_12_2_CSV.zip"
QUALITY_URL = "https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp/Developed_6_Portfolios_ME_OP_CSV.zip"
AQR_URL = "https://www.aqr.com/-/media/AQR/Documents/Insights/Data-Sets/Time-Series-Momentum-Factors-Monthly.xlsx"


def archive(member: str, text: str) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as zipped:
        zipped.writestr(member, text)
    return output.getvalue()


def frozen(directory: Path, name: str, content: bytes, url: str) -> SourceArtifact:
    path = directory / name
    path.write_bytes(content)
    return SourceArtifact(url, NOW, "a" * 64, path, len(content), "test")


def style_archive(member: str, columns: str, rows: str, table: str = "Average Value Weighted Returns -- Monthly") -> bytes:
    return archive(
        member,
        f"Missing data are indicated by -99.99.\n\n  {table}\n,{columns}\n{rows}\n\n  Average Equal Weighted Returns -- Monthly\n,{columns}\n",
    )


def aqr_workbook(
    rows: tuple[tuple[object, object], ...],
    *,
    column: str = "TSMOM",
    number_format: str = "0.00%",
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TSMOM Factors"
    sheet["A1"] = "AQR Capital Management, LLC - Time Series Momentum: Factors, Monthly"
    sheet["A3"] = "This file contains the excess returns of the long/short Time Series Momentum (TSMOM) factors."
    sheet["A7"] = "We report the monthly excess returns of the time series momentum factors for all assets (no superscript), global equity indices (EQ), currencies (FX), fixed income (FI), and commodities (CM)"
    for cell, value in zip(("B18", "C18", "D18", "E18", "F18"), (column, "TSMOM^CM", "TSMOM^EQ", "TSMOM^FI", "TSMOM^FX")):
        sheet[cell] = value
    for row_index, (month, value) in enumerate(rows, 19):
        sheet.cell(row_index, 1, month)
        sheet.cell(row_index, 2, value).number_format = number_format
        for column_index in range(3, 7):
            sheet.cell(row_index, column_index, Decimal("0.01")).number_format = "0.00%"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def inputs(directory: Path, **changes: object) -> PublicWorldInputs:
    months: list[str] = []
    current = date(1977, 1, 1)
    while current <= date(1990, 6, 1):
        months.append(f"{current.year}{current.month:02d}")
        current = date(current.year + (current.month == 12), current.month % 12 + 1, 1)
    international_rows = "\n".join(f"{month} 5.00" for month in months)
    us_rows = "\n".join(f"{month},5.00,0.10" for month in months)
    international = archive(
        "Ind_all.Dat",
        f"Value-Weight Dollar Returns\nDate Mkt\n{international_rows}\n",
    )
    us = archive(
        "F-F_Research_Data_Factors.csv",
        f"Date,Mkt-RF,RF\n{us_rows}\n",
    )
    developed = archive("Developed_3_Factors.csv", "Date,Mkt-RF,RF\n199007,7.00,0.20\n")
    fields: dict[str, object] = {
        "international_archive": frozen(directory, "international.zip", international, "https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp/F-F_International_Indices.zip"),
        "us_archive": frozen(directory, "us.zip", us, "https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp/F-F_Research_Data_Factors_CSV.zip"),
        "developed_archive": frozen(directory, "developed.zip", developed, "https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp/Developed_3_Factors_CSV.zip"),
        "wdi_caps": {"USA": {year: Decimal("40") for year in range(1975, 1989)}, "WLD": {year: Decimal("100") for year in range(1975, 1989)}},
        "wdi_source_url": "https://api.worldbank.org/v2",
        "wdi_source_hash": "b" * 64,
        "retrieved_at": NOW,
    }
    fields.update(changes)
    return PublicWorldInputs(**fields)  # type: ignore[arg-type]


def fx_row(day: date, value: str, series: str = "DEXUSEU") -> ObservationRow:
    return ObservationRow(series, day, day, NOW, Decimal(value), "USD_per_EUR", "https://fred.stlouisfed.org/graph/fredgraph.csv", NOW, "c" * 64)


def month_end(year: int, month: int) -> date:
    first_next = date(year + (month == 12), month % 12 + 1, 1)
    return first_next - timedelta(days=1)


class MarketProxyTests(TestCase):
    def test_exact_academic_sleeves_are_normalized_without_rescaling_aqr(self):
        momentum_columns = "SMALL LoPRIOR,ME1 PRIOR2,SMALL HiPRIOR,BIG LoPRIOR,ME2 PRIOR2,BIG HiPRIOR"
        quality_columns = "SMALL LoOP,ME1 OP2,SMALL HiOP,BIG LoOP,ME2 OP2,BIG HiOP"
        fixtures = (
            (
                "french_developed_momentum_archive",
                MOMENTUM_URL,
                "momentum.zip",
                style_archive(
                    "Developed_6_Portfolios_ME_Prior_12_2.csv",
                    momentum_columns,
                    "200601,1,2,3,4,5,6\n200602,2,3,4,5,6,7",
                ),
                "FRENCH_DEVELOPED_BIG_HIPRIOR_USD",
                (Decimal("0.06"), Decimal("0.07")),
            ),
            (
                "french_developed_quality_archive",
                QUALITY_URL,
                "quality.zip",
                style_archive(
                    "Developed_6_Portfolios_ME_OP.csv",
                    quality_columns,
                    "200601,1,2,3,4,5,6\n200602,-1,0,1,2,3,4",
                ),
                "FRENCH_DEVELOPED_BIG_ROBUST_USD",
                (Decimal("0.06"), Decimal("0.04")),
            ),
            (
                "french_aqr_tsmom_archive",
                AQR_URL,
                "tsmom.xlsx",
                aqr_workbook(
                    (
                        (datetime(2006, 1, 31), Decimal("0.03")),
                        (datetime(2006, 2, 28), Decimal("-0.01")),
                    )
                ),
                "AQR_TSMOM_EXCESS_RETURN_USD",
                (Decimal("0.03"), Decimal("-0.01")),
            ),
        )
        with TemporaryDirectory() as directory:
            root = Path(directory)
            for parser_kind, url, name, content, series_id, expected in fixtures:
                with self.subTest(parser_kind=parser_kind):
                    rows = parse_french_archive(frozen(root, name, content, url), parser_kind)
                    self.assertEqual(tuple(row.series_id for row in rows), (series_id, series_id))
                    self.assertEqual(tuple(row.value for row in rows), expected)
                    self.assertEqual(tuple(row.observation_date for row in rows), (date(2006, 1, 31), date(2006, 2, 28)))
                    self.assertTrue(all(row.unit == "ratio" for row in rows))

    def test_academic_sleeves_reject_wrong_table_columns_units_and_sentinels(self):
        momentum_columns = "SMALL LoPRIOR,ME1 PRIOR2,SMALL HiPRIOR,BIG LoPRIOR,ME2 PRIOR2,BIG HiPRIOR"
        failures = (
            style_archive(
                "Developed_6_Portfolios_ME_Prior_12_2.csv",
                momentum_columns,
                "200601,1,2,3,4,5,6",
                table="Average Equal Weighted Returns -- Monthly",
            ),
            style_archive(
                "Developed_6_Portfolios_ME_Prior_12_2.csv",
                momentum_columns.replace("BIG HiPRIOR", "BIG Wrong"),
                "200601,1,2,3,4,5,6",
            ),
            style_archive(
                "Developed_6_Portfolios_ME_Prior_12_2.csv",
                momentum_columns,
                "200601,1,2,3,4,5,-99.99",
            ),
        )
        with TemporaryDirectory() as directory:
            root = Path(directory)
            for index, content in enumerate(failures):
                with self.subTest(index=index), self.assertRaises(ValueError):
                    parse_french_archive(
                        frozen(root, f"bad-{index}.zip", content, MOMENTUM_URL),
                        "french_developed_momentum_archive",
                    )
            for index, content in enumerate(
                (
                    aqr_workbook(((datetime(2006, 1, 31), Decimal("0.03")),), column="WRONG"),
                    aqr_workbook(((datetime(2006, 1, 31), Decimal("0.03")),), number_format="0.00"),
                    aqr_workbook(((datetime(2006, 1, 31), Decimal("-99.99")),)),
                )
            ):
                with self.subTest(aqr=index), self.assertRaises(ValueError):
                    parse_french_archive(
                        frozen(root, f"bad-{index}.xlsx", content, AQR_URL),
                        "french_aqr_tsmom_archive",
                    )

    def test_academic_sleeves_reject_duplicates_gaps_and_post_start_parse_errors(self):
        columns = "SMALL LoPRIOR,ME1 PRIOR2,SMALL HiPRIOR,BIG LoPRIOR,ME2 PRIOR2,BIG HiPRIOR"
        french_failures = (
            "200601,1,2,3,4,5,6\n200601,1,2,3,4,5,6",
            "200601,1,2,3,4,5,6\n200603,1,2,3,4,5,6",
            "200601,1,2,3,4,5,6\nBROKEN,1,2,3,4,5,6\n200602,1,2,3,4,5,6",
        )
        aqr_failures = (
            ((datetime(2006, 1, 31), Decimal("0.03")), (datetime(2006, 1, 31), Decimal("0.04"))),
            ((datetime(2006, 1, 31), Decimal("0.03")), (datetime(2006, 3, 31), Decimal("0.04"))),
            ((datetime(2006, 1, 31), Decimal("0.03")), ("BROKEN", Decimal("0.04")), (datetime(2006, 2, 28), Decimal("0.05"))),
        )
        with TemporaryDirectory() as directory:
            root = Path(directory)
            for index, rows in enumerate(french_failures):
                with self.subTest(french=index), self.assertRaises(ValueError):
                    parse_french_archive(
                        frozen(
                            root,
                            f"bad-{index}.zip",
                            style_archive("Developed_6_Portfolios_ME_Prior_12_2.csv", columns, rows),
                            MOMENTUM_URL,
                        ),
                        "french_developed_momentum_archive",
                    )
            for index, rows in enumerate(aqr_failures):
                with self.subTest(aqr=index), self.assertRaises(ValueError):
                    parse_french_archive(
                        frozen(root, f"bad-{index}.xlsx", aqr_workbook(rows), AQR_URL),
                        "french_aqr_tsmom_archive",
                    )

    def test_fx_direction_sentinel(self):
        self.assertAlmostEqual(convert_usd_to_eur(0.10, 1.0, 0.9), -0.01, places=12)

    def test_splice_is_exact_and_label_is_not_msci(self):
        with TemporaryDirectory() as directory:
            result = build_public_world_monthly(inputs(Path(directory)))
        by_month = {item.observation.observation_date: item for item in result}
        self.assertEqual(by_month[date(1990, 6, 30)].segment, "RECONSTRUCTED_SPLICE")
        self.assertEqual(by_month[date(1990, 7, 31)].segment, "FF_DEVELOPED")
        self.assertEqual(by_month[date(1990, 7, 31)].observation.series_id, CANONICAL_WORLD_LABEL)
        self.assertNotIn("MSCI", by_month[date(1990, 7, 31)].observation.series_id)

    def test_derived_rows_preserve_all_input_provenance_and_use_composite_hash(self):
        with TemporaryDirectory() as directory:
            result = build_public_world_monthly(inputs(Path(directory)))
        reconstructed = next(item for item in result if item.observation.observation_date == date(1990, 6, 30))
        developed = next(item for item in result if item.observation.observation_date == date(1990, 7, 31))
        self.assertEqual(
            {source.source_url for source in reconstructed.input_provenance},
            {
                "https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp/F-F_International_Indices.zip",
                "https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp/F-F_Research_Data_Factors_CSV.zip",
                "https://api.worldbank.org/v2",
            },
        )
        self.assertEqual([source.source_url for source in developed.input_provenance], ["https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp/Developed_3_Factors_CSV.zip"])
        self.assertNotEqual(reconstructed.observation.source_hash, "a" * 64)
        self.assertNotEqual(reconstructed.observation.source_url, reconstructed.input_provenance[0].source_url)
        self.assertEqual(reconstructed.observation.retrieved_at, NOW)

    def test_monthly_outcome_becomes_available_only_at_month_end(self):
        with TemporaryDirectory() as directory:
            result = build_public_world_monthly(inputs(Path(directory)))
        june = next(item.observation for item in result if item.observation.observation_date == date(1990, 6, 30))
        self.assertEqual(june.available_at, datetime(1990, 6, 30, 23, 59, 59, tzinfo=timezone.utc))
        self.assertIsNone(asof_select((june,), datetime(1990, 6, 30, 23, 59, 58, tzinfo=timezone.utc)))
        self.assertEqual(asof_select((june,), datetime(1990, 7, 1, tzinfo=timezone.utc)), june)

    def test_warm_up_starts_january_1977_and_reported_months_have_no_gaps_or_duplicates(self):
        with TemporaryDirectory() as directory:
            result = build_public_world_monthly(inputs(Path(directory)))
        months = [item.observation.observation_date for item in result]
        self.assertEqual(months[:2], [date(1977, 1, 31), date(1977, 2, 28)])
        self.assertEqual(len(months), len(set(months)))
        reported = months[months.index(date(1979, 1, 31)):]
        expected = date(1979, 1, 31)
        for month in reported:
            self.assertEqual(month, expected)
            expected = date(expected.year + (expected.month == 12), expected.month % 12 + 1, 1)
            expected = date(expected.year, expected.month, 1)
            expected = date(expected.year + (expected.month == 12), expected.month % 12 + 1, 1).fromordinal(date(expected.year + (expected.month == 12), expected.month % 12 + 1, 1).toordinal() - 1)

    def test_rejects_wrong_member_table_columns_and_french_sentinel(self):
        variants = (
            ("wrong member", {"international_archive": None}),
            ("wrong table", {"international_archive": "not a table"}),
            ("wrong columns", {"international_archive": "wrong columns"}),
            ("sentinel", {"international_archive": "sentinel"}),
        )
        for label, change in variants:
            with self.subTest(label=label), TemporaryDirectory() as directory:
                base = inputs(Path(directory))
                if change["international_archive"] is None:
                    bad = frozen(Path(directory), "wrong.zip", archive("wrong.Dat", "x"), base.international_archive.source_url)
                    supplied = PublicWorldInputs(bad, base.us_archive, base.developed_archive, base.wdi_caps, base.wdi_source_url, base.wdi_source_hash, NOW)
                elif change["international_archive"]:
                    text = "Other Table\nDate Mkt\n197701 1.00\n" if change["international_archive"] == "not a table" else "Value-Weight Dollar Returns\nDate Wrong\n197701 1.00\n" if change["international_archive"] == "wrong columns" else "Value-Weight Dollar Returns\nDate Mkt\n197701 -99.99\n"
                    bad = frozen(Path(directory), "bad.zip", archive("Ind_all.Dat", text), base.international_archive.source_url)
                    supplied = PublicWorldInputs(bad, base.us_archive, base.developed_archive, base.wdi_caps, base.wdi_source_url, base.wdi_source_hash, NOW)
                with self.assertRaises(ValueError):
                    build_public_world_monthly(supplied)

    def test_french_percent_returns_are_divided_by_one_hundred_once(self):
        with TemporaryDirectory() as directory:
            result = build_public_world_monthly(inputs(Path(directory)))
        july = next(row for row in result if row.observation.observation_date == date(1990, 7, 31))
        self.assertAlmostEqual(float(july.observation.value), 0.072, places=12)

    def test_wdi_uses_mandatory_y_minus_two_and_rejects_out_of_bounds_weight(self):
        with TemporaryDirectory() as directory:
            caps = {aggregate: dict(values) for aggregate, values in inputs(Path(directory)).wdi_caps.items()}
            caps["USA"][1983] = Decimal("50")
            base = inputs(Path(directory), wdi_caps=caps)
            result = build_public_world_monthly(base)
            jan_1985 = next(row for row in result if row.observation.observation_date == date(1985, 1, 31))
            self.assertAlmostEqual(float(jan_1985.observation.value), 0.0505, places=12)
            missing = {"USA": {1977: Decimal("40")}, "WLD": {1977: Decimal("100")}}
            with self.assertRaises(ValueError):
                build_public_world_monthly(PublicWorldInputs(base.international_archive, base.us_archive, base.developed_archive, missing, base.wdi_source_url, base.wdi_source_hash, NOW))
            invalid = {"USA": {**base.wdi_caps["USA"], 1975: Decimal("101")}, "WLD": base.wdi_caps["WLD"]}
            with self.assertRaises(ValueError):
                build_public_world_monthly(PublicWorldInputs(base.international_archive, base.us_archive, base.developed_archive, invalid, base.wdi_source_url, base.wdi_source_hash, NOW))

    def test_fx_quote_direction_december_sentinel_month_end_and_daily_carry_boundary(self):
        normalized = normalize_fx_monthly((fx_row(date(2023, 12, 1), "1.100"), fx_row(date(2023, 12, 29), "1.105")))
        self.assertEqual(normalized[0].raw_quote, "USD_per_EUR")
        self.assertAlmostEqual(normalized[0].raw, 1.105, places=12)
        self.assertAlmostEqual(normalized[0].eur_per_usd, 1 / 1.105, places=12)
        self.assertEqual(normalized[0].observation.unit, "EUR_per_USD")
        self.assertEqual(normalized[0].observation.source_hash, "c" * 64)
        carried = carry_fx_daily(normalized, (date(2024, 1, 5),))
        self.assertAlmostEqual(carried[date(2024, 1, 5)].value, 1 / 1.105, places=12)
        self.assertEqual(carried[date(2024, 1, 5)].fixing.observation.source_hash, "c" * 64)
        with self.assertRaises(ValueError):
            carry_fx_daily(normalized, (date(2024, 1, 9),))

    def test_fx_month_end_selection_is_total_and_independent_of_input_order(self):
        earlier = fx_row(date(2024, 1, 31), "1.100")
        later = ObservationRow("DEXUSEU", date(2024, 1, 31), date(2024, 1, 31), NOW, Decimal("1.200"), "USD_per_EUR", earlier.source_url, NOW + timedelta(seconds=1), "d" * 64)
        forward = normalize_fx_monthly((earlier, later))
        backward = normalize_fx_monthly((later, earlier))
        self.assertEqual(forward, backward)
        self.assertEqual(repr(forward).encode("utf-8"), repr(backward).encode("utf-8"))
        self.assertAlmostEqual(forward[0].raw, 1.2, places=12)

    def test_validation_metrics_minimum_sample_missing_nav_and_threshold_failure(self):
        starts = date(2019, 1, 31)
        proxy = {date(2019 + (index // 12), index % 12 + 1, 28): 0.01 + (index % 3) * 0.001 for index in range(60)}
        nav = dict(proxy)
        passed = validate_world_proxy(proxy, nav)
        self.assertEqual(passed.n, 60)
        self.assertEqual(passed.status, "PASS")
        self.assertAlmostEqual(passed.correlation, 1.0, places=12)
        self.assertAlmostEqual(passed.beta, 1.0, places=12)
        self.assertAlmostEqual(passed.tracking_error, 0.0, places=12)
        self.assertAlmostEqual(passed.cagr_gap, 0.0, places=12)
        self.assertAlmostEqual(passed.max_drawdown_gap, 0.0, places=12)
        self.assertEqual(validate_world_proxy(proxy, None).status, "FAIL_MISSING_OFFICIAL_NAV")
        too_short = dict(list(nav.items())[:-1])
        self.assertEqual(validate_world_proxy(proxy, too_short).status, "FAIL_INSUFFICIENT_NAV_SAMPLE")
        missing_middle = dict(nav)
        del missing_middle[date(2021, 6, 28)]
        self.assertEqual(validate_world_proxy(proxy, missing_middle).status, "FAIL_INSUFFICIENT_NAV_SAMPLE")
        nonfinite = dict(nav)
        nonfinite[date(2021, 6, 28)] = float("nan")
        self.assertEqual(validate_world_proxy(proxy, nonfinite).status, "FAIL_INSUFFICIENT_NAV_SAMPLE")
        failed = validate_world_proxy(proxy, {month: -value for month, value in proxy.items()})
        self.assertEqual(failed.status, "FAIL")

    def test_validation_matches_explicit_ols_te_cagr_and_drawdown_formulas(self):
        nav = {month_end(2019 + index // 12, index % 12 + 1): (-0.04 if index % 11 == 0 else 0.01 + (index % 5) * 0.003) for index in range(60)}
        proxy = {month: 0.002 + 0.97 * value + (0.001 if index % 2 else -0.001) for index, (month, value) in enumerate(nav.items())}
        metrics = validate_world_proxy(proxy, nav)
        p, n = list(proxy.values()), list(nav.values())
        mean_p, mean_n = sum(p) / 60, sum(n) / 60
        covariance = sum((left - mean_p) * (right - mean_n) for left, right in zip(p, n)) / 59
        variance_n = sum((value - mean_n) ** 2 for value in n) / 59
        differences = [left - right for left, right in zip(p, n)]
        expected_beta = covariance / variance_n
        expected_te = math.sqrt(sum((value - sum(differences) / 60) ** 2 for value in differences) / 59) * math.sqrt(12)
        expected_cagr_gap = abs(math.prod(1 + value for value in p) ** (12 / 60) - 1 - (math.prod(1 + value for value in n) ** (12 / 60) - 1))
        def drawdown(values: list[float]) -> float:
            wealth = peak = 1.0
            lowest = 0.0
            for value in values:
                wealth *= 1 + value
                peak = max(peak, wealth)
                lowest = min(lowest, wealth / peak - 1)
            return lowest
        self.assertAlmostEqual(metrics.beta, expected_beta, places=12)
        self.assertAlmostEqual(metrics.tracking_error, expected_te, places=12)
        self.assertAlmostEqual(metrics.cagr_gap, expected_cagr_gap, places=12)
        self.assertAlmostEqual(metrics.max_drawdown_gap, abs(drawdown(p) - drawdown(n)), places=12)
